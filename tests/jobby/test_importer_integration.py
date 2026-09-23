"""Conservative, offline legacy-import reconciliation tests."""

from __future__ import annotations

import logging
import stat
import sys
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy import func, select

from jobby.analytics import score_calibration
from jobby.config import JobbyPaths
from jobby.db import Database
from jobby.enums import ApplicationStage, ImportReviewStatus, JobStatus
from jobby.importer import (
    LegacyImporter,
    _normalize_loaded_data,
    discover_legacy_artifacts,
    extract_document_text,
    sha256_file,
    status_to_stage,
)
from jobby.models import (
    Application,
    Artifact,
    AuditEvent,
    Company,
    DocumentVersion,
    Evaluation,
    ImportReview,
    Job,
    LegacyMetric,
    LegacySeenIdentifier,
    SourceObservation,
    StageEvent,
)


def _paths(root: Path) -> JobbyPaths:
    home = root / "jobby-home"
    data = home / "data"
    config = home / "config"
    cache = home / "cache"
    return JobbyPaths(
        data_dir=data,
        config_dir=config,
        cache_dir=cache,
        database=data / "jobby.sqlite3",
        artifacts_dir=data / "artifacts",
        backups_dir=data / "backups",
        logs_dir=data / "logs",
        config_file=config / "config.toml",
    ).ensure()


def _importer(
    tmp_path: Path, workspace: Path | None = None
) -> tuple[Database, LegacyImporter, Path]:
    workspace = workspace or tmp_path / "legacy"
    workspace.mkdir(parents=True, exist_ok=True)
    paths = _paths(tmp_path)
    database = Database(paths=paths)
    return database, LegacyImporter(database, paths=paths), workspace


def _write(path: Path, content: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    return path


def test_structured_import_rejects_non_json_numbers_and_key_collisions(
    tmp_path: Path,
) -> None:
    with pytest.raises(ValueError, match="non-finite"):
        _normalize_loaded_data({"value": float("nan")})
    with pytest.raises(ValueError, match="collide"):
        _normalize_loaded_data({1: "numeric", "1": "text"})

    source = _write(tmp_path / "source.txt", "content")
    with pytest.raises(TypeError, match="integer"):
        sha256_file(source, True)
    with pytest.raises(ValueError, match="positive"):
        sha256_file(source, 0)


def test_import_preserves_private_url_as_observation_only(tmp_path: Path) -> None:
    database, importer, workspace = _importer(tmp_path)
    private_url = "http://127.0.0.1:8080/internal-job"
    _write(
        workspace / "new_jobs_2026-07-14.md",
        "## Internal Counsel\n"
        "**Example Co** | GH\n\n"
        "Keywords matched: policy\n"
        f"{private_url}\n\n---\n",
    )

    importer.run(workspace)

    with database.session() as session:
        job = session.scalar(select(Job).where(Job.title == "Internal Counsel"))
        observation = session.scalar(
            select(SourceObservation).where(SourceObservation.source_url == private_url)
        )
        assert job is not None
        assert job.launch_url is None
        assert job.comparison_url is None
        assert job.canonical_url is None
        assert observation is not None
    database.dispose()


def _pipeline_line(score: str = "4.2") -> str:
    return (
        "# Evaluation Pipeline\n\n"
        f"- [{score}/5] Example Co — Policy Analyst | PINNED — strong fit | "
        "[Report](reports/example.md)\n"
    )


def _scan_report(title: str = "Policy Analyst", *, include_url: bool = True) -> str:
    url = (
        "\nhttps://boards.greenhouse.io/example/jobs/123?gh_jid=123\n"
        if include_url
        else "\n"
    )
    return (
        f"## {title}\n"
        "**Example Co** | GH\n\n"
        "Keywords matched: policy, governance\n"
        f"{url}\n---\n"
    )


def _tracker(entries: list[tuple[str, str, str]]) -> str:
    blocks = ["# Application Tracker"]
    for index, (company, title, status_text) in enumerate(entries, 1):
        blocks.append(
            f"""
### {index}.1 {company} — {title}
| Field | Details |
|-------|---------|
| **Status** | {status_text} |
| **Skills Match** | **8/10** — strong |
| **Application Success** | **5/10** — plausible |
| **Link** | [Apply](https://example.test/{index}) |
""".strip()
        )
    return "\n\n".join(blocks) + "\n"


def _table_counts(database: Database) -> dict[str, int]:
    models = (
        Job,
        SourceObservation,
        Evaluation,
        LegacyMetric,
        Application,
        StageEvent,
        ImportReview,
        DocumentVersion,
        LegacySeenIdentifier,
    )
    with database.session() as session:
        return {
            model.__tablename__: session.scalar(select(func.count()).select_from(model))
            for model in models
        }


def test_discovery_of_recognized_directories_is_independent_of_process_cwd(
    tmp_path, monkeypatch
):
    workspace = tmp_path / "detached-workspace"
    expected = {
        "reports/evaluation.md",
        "jds/role.md",
        "output/tailored-resume.html",
        "interview-prep/notes.md",
        "hawaii-ag-applications/form.pdf",
    }
    for relative in expected:
        path = workspace / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"source")
    other_cwd = tmp_path / "unrelated-cwd"
    other_cwd.mkdir()
    monkeypatch.chdir(other_cwd)

    discovered = {
        path.relative_to(workspace).as_posix()
        for path in discover_legacy_artifacts(workspace)
    }
    assert discovered == expected


def test_repeat_import_is_stable_and_never_rewrites_workspace_sources(tmp_path):
    database, importer, workspace = _importer(tmp_path)
    _write(workspace / "data/pipeline.md", _pipeline_line())
    _write(workspace / "new_jobs_2026-03-01.md", _scan_report())
    _write(
        workspace / "Application Tracker - Test.md",
        _tracker(
            [("Tracked Co", "Applied Researcher", "✅ **SUBMITTED — March 3, 2026**")]
        ),
    )
    _write(
        workspace / "Fixture Candidate Resume 2026.md",
        "# Fixture Candidate\n\nOriginal resume text.\n",
    )
    _write(workspace / "job_monitor_state.json", '{"seen_jobs": ["gh:1", "gh:2"]}')
    source_state = {
        path.relative_to(workspace).as_posix(): (
            sha256_file(path),
            path.stat().st_mtime_ns,
        )
        for path in workspace.rglob("*")
        if path.is_file()
    }

    first = importer.run(workspace)
    first_counts = _table_counts(database)
    second = importer.run(workspace)
    second_counts = _table_counts(database)

    assert first.imported_artifacts == len(source_state)
    assert second.imported_artifacts == 0
    assert second.skipped_artifacts == len(source_state)
    assert first_counts == second_counts
    assert not any(issue.source_path == "data/pipeline.md" for issue in first.unparsed)
    assert {
        path.relative_to(workspace).as_posix(): (
            sha256_file(path),
            path.stat().st_mtime_ns,
        )
        for path in workspace.rglob("*")
        if path.is_file()
    } == source_state
    with database.session() as session:
        # Each attempted import is itself retained in the append-only activity log.
        assert (
            session.scalar(
                select(func.count())
                .select_from(AuditEvent)
                .where(AuditEvent.action == "workspace.imported")
            )
            == 2
        )
        assert session.scalar(
            select(func.count())
            .select_from(Artifact)
            .where(Artifact.source_immutable.is_(True))
        ) == len(source_state)
        exports = list(
            session.scalars(
                select(Artifact).where(Artifact.source_immutable.is_(False))
            )
        )
        assert len(exports) == 4
        assert all(artifact.kind.value == "generated_export" for artifact in exports)
        assert (
            session.scalar(
                select(func.count())
                .select_from(AuditEvent)
                .where(AuditEvent.action == "import.reconciliation_exported")
            )
            == 2
        )
    database.dispose()


def test_changed_source_and_resume_create_immutable_artifact_and_document_versions(
    tmp_path,
):
    database, importer, workspace = _importer(tmp_path)
    resume = _write(
        workspace / "Fixture Candidate Resume 2026.md",
        "# Fixture Candidate\n\nVersion one.\n",
    )
    first_source_hash = sha256_file(resume)
    importer.run(workspace)

    resume.write_text(
        "# Fixture Candidate\n\nVersion two with a new fact.\n", encoding="utf-8"
    )
    second_source_hash = sha256_file(resume)
    importer.run(workspace)

    with database.session() as session:
        artifacts = list(
            session.scalars(
                select(Artifact)
                .where(Artifact.source_path == resume.name)
                .order_by(Artifact.created_at)
            )
        )
        documents = list(
            session.scalars(select(DocumentVersion).order_by(DocumentVersion.version))
        )
    assert [artifact.content_hash for artifact in artifacts] == [
        first_source_hash,
        second_source_hash,
    ]
    assert [document.version for document in documents] == [1, 2]
    assert [document.content_markdown for document in documents] == [
        "# Fixture Candidate\n\nVersion one.\n",
        "# Fixture Candidate\n\nVersion two with a new fact.\n",
    ]
    assert all(artifact.source_immutable for artifact in artifacts)
    assert all(Path(artifact.stored_path).read_bytes() for artifact in artifacts)
    assert all(
        Path(artifact.stored_path).stat().st_mode
        & (stat.S_IWUSR | stat.S_IWGRP | stat.S_IWOTH)
        == 0
        for artifact in artifacts
    )
    database.dispose()


def test_resume_formats_collapse_into_three_semantic_variants_with_one_canonical_master(
    tmp_path,
):
    from docx import Document

    database, importer, workspace = _importer(tmp_path)
    master_text = (
        "Fixture Candidate legal technology policy governance patent portfolio "
        "New Media Rights Clinic artificial intelligence research"
    )
    _write(
        workspace / "Fixture Candidate Resume 2026.md",
        f"# Fixture Candidate\n\n{master_text}\n",
    )
    master_docx = Document()
    master_docx.add_heading("Fixture Candidate", level=1)
    master_docx.add_paragraph(master_text)
    master_docx.save(workspace / "Fixture Candidate Resume 2026.docx")

    _write(
        workspace / "Fixture_Candidate_Anthropic_Tailored_Resume.md",
        "# Fixture Candidate\n\nAnthropic safeguards enforcement and AI evaluation systems.\n",
    )
    legora_text = (
        "Fixture Candidate legal data analyst structured datasets prompt engineering "
        "quality assurance legal workflows"
    )
    _write(
        workspace / "output/legora-legal-data-analyst-resume.md",
        f"# Fixture Candidate\n\n{legora_text}\n",
    )
    _write(
        workspace / "output/legora-legal-data-analyst-resume.html",
        f"<html><body><h1>Fixture Candidate</h1><p>{legora_text}</p></body></html>",
    )

    importer.run(workspace)

    with database.session() as session:
        documents = list(
            session.scalars(
                select(DocumentVersion)
                .where(DocumentVersion.kind == "resume")
                .order_by(DocumentVersion.name)
            )
        )
    assert len(documents) == 3
    canonical = [document for document in documents if document.is_canonical]
    assert len(canonical) == 1
    assert canonical[0].approval_state.value == "approved"
    assert len(canonical[0].provenance) == 1
    with database.session() as session:
        linked_sources = session.scalar(
            select(func.count())
            .select_from(Artifact)
            .where(Artifact.document_version_id == canonical[0].id)
        )
    assert linked_sources == 2
    assert canonical[0].content_markdown.startswith("# Fixture Candidate")
    assert (
        sum(document.approval_state.value == "approved" for document in documents) == 1
    )
    database.dispose()


def test_malformed_markdown_is_queued_once_while_valid_pipeline_rows_are_not(tmp_path):
    database, importer, workspace = _importer(tmp_path)
    _write(
        workspace / "data/pipeline.md",
        _pipeline_line()
        + "- [not-a-score/5] Broken entry without the required structure\n",
    )
    _write(workspace / "new_jobs_2026-03-02.md", _scan_report(include_url=False))

    first = importer.run(workspace)
    importer.run(workspace)

    assert {(issue.source_path, issue.record_key) for issue in first.unparsed} == {
        ("data/pipeline.md", "line:4"),
        ("new_jobs_2026-03-02.md", "block:1"),
    }
    with database.session() as session:
        reviews = list(session.scalars(select(ImportReview)))
    assert {(review.source_path, review.record_key) for review in reviews} == {
        ("data/pipeline.md", "line:4"),
        ("new_jobs_2026-03-02.md", "block:1"),
    }
    database.dispose()


def test_source_state_versions_are_unioned_without_inventing_jobs(tmp_path):
    database, importer, workspace = _importer(tmp_path)
    _write(
        workspace / "job_monitor_state.json",
        '{"seen_jobs": ["greenhouse:a", "lever:b"], "last_run": "2026-03-03"}',
    )
    _write(
        workspace / "job_monitor_state.json.bak-2026-03-02",
        '{"seen_jobs": ["lever:b", "usajobs:c"], "last_run": "2026-03-02"}',
    )

    report = importer.run(workspace)
    importer.run(workspace)

    with database.session() as session:
        identifiers = set(session.scalars(select(LegacySeenIdentifier.source_uid)))
        jobs = session.scalar(select(func.count()).select_from(Job))
        review = session.scalar(select(ImportReview))
    assert identifiers == {"greenhouse:a", "lever:b", "usajobs:c"}
    assert jobs == 0
    assert review is not None and review.status == ImportReviewStatus.RESOLVED
    assert report.conflicts == []
    assert len(report.resolved) == 1
    assert "unioned without inventing job records" in report.resolved[0].message
    database.dispose()


def test_multi_role_report_is_preserved_and_dismissed_without_recurring_review(
    tmp_path,
):
    database, importer, workspace = _importer(tmp_path)
    _write(
        workspace / "reports/multi-role-analysis.md",
        "# Multi-role analysis\n\n"
        "## 1. First role\nhttps://example.test/jobs/one\n\n"
        "## 2. Second role\nhttps://example.test/jobs/two\n",
    )

    first = importer.run(workspace)
    second = importer.run(workspace)

    with database.session() as session:
        reviews = list(session.scalars(select(ImportReview)))
    assert first.unparsed == second.unparsed == []
    assert len(first.resolved) == len(second.resolved) == 1
    assert len(reviews) == 1
    assert reviews[0].status == ImportReviewStatus.DISMISSED
    database.dispose()


def test_saved_job_description_links_by_exact_canonical_url_and_populates_job(
    tmp_path,
):
    database, importer, workspace = _importer(tmp_path)
    database.initialize()
    url = "https://boards.example.test/acme/jobs/123"
    with database.session() as session:
        company = Company(name="Acme", normalized_name="acme")
        session.add(company)
        session.flush()
        job = Job(
            company_id=company.id,
            title="Policy Counsel",
            normalized_title="policy counsel",
            canonical_url=url,
        )
        session.add(job)
        session.flush()
        job_id = job.id
    _write(
        workspace / "jds/acme-policy-counsel.md",
        "# JD: Acme — Policy Counsel\n\n"
        f"**URL:** {url}\n\n"
        "Advise the company on AI policy.\n",
    )

    report = importer.run(workspace)

    with database.session() as session:
        job = session.get(Job, job_id)
        artifact = session.scalar(
            select(Artifact).where(Artifact.source_path == "jds/acme-policy-counsel.md")
        )
        review = session.scalar(
            select(ImportReview).where(
                ImportReview.record_key == "job_description_link"
            )
        )
    assert report.unparsed == []
    assert job is not None and "Advise the company" in (job.description or "")
    assert job.description_hash
    assert artifact is not None and artifact.job_id == job_id
    assert review is not None and review.status == ImportReviewStatus.RESOLVED
    assert review.proposed_json["match_method"] == "canonical_url"
    database.dispose()


def test_saved_job_description_links_by_unique_heading_in_either_orientation(
    tmp_path,
):
    database, importer, workspace = _importer(tmp_path)
    database.initialize()
    with database.session() as session:
        company = Company(name="Broadcom", normalized_name="broadcom")
        session.add(company)
        session.flush()
        job = Job(
            company_id=company.id,
            title="AI Innovation Specialist, Legal",
            normalized_title="ai innovation specialist legal",
        )
        session.add(job)
        session.flush()
        job_id = job.id
    _write(
        workspace / "jds/broadcom-ai-innovation-specialist-legal.md",
        "# AI Innovation Specialist - Legal — Broadcom\n\n"
        "**Req ID:** R025347\n\nBuild legal automation systems.\n",
    )

    importer.run(workspace)

    with database.session() as session:
        artifact = session.scalar(
            select(Artifact).where(
                Artifact.source_path == "jds/broadcom-ai-innovation-specialist-legal.md"
            )
        )
        review = session.scalar(select(ImportReview))
    assert artifact is not None and artifact.job_id == job_id
    assert review is not None and review.status == ImportReviewStatus.RESOLVED
    assert review.proposed_json["match_method"] == "normalized_company_title"
    database.dispose()


def test_repeated_scan_reports_keep_each_observation_and_title_snapshot_without_liveness_inference(
    tmp_path,
):
    database, importer, workspace = _importer(tmp_path)
    original = _scan_report("Policy Analyst")
    _write(workspace / "new_jobs_2026-03-01.md", original)
    # An exact copied report is still a distinct dated source observation.
    _write(workspace / "new_jobs_2026-03-02.md", original)
    _write(workspace / "new_jobs_2026-03-03.md", _scan_report("Senior Policy Analyst"))

    importer.run(workspace)
    importer.run(workspace)

    with database.session() as session:
        jobs = list(session.scalars(select(Job)))
        observations = list(
            session.scalars(
                select(SourceObservation).order_by(SourceObservation.observed_at)
            )
        )
    assert len(jobs) == 1
    assert [item.title_snapshot for item in observations] == [
        "Policy Analyst",
        "Policy Analyst",
        "Senior Policy Analyst",
    ]
    assert [item.observed_at.date() for item in observations] == [
        date(2026, 3, 1),
        date(2026, 3, 2),
        date(2026, 3, 3),
    ]
    assert all(item.is_live is None for item in observations)
    assert len({item.import_key for item in observations}) == 3
    job = jobs[0]
    assert job.title == "Policy Analyst"
    assert job.liveness_known is False
    assert job.consecutive_misses == 0
    assert job.explicit_closure is False
    assert job.closed_at is None
    assert job.status == JobStatus.DISCOVERED
    database.dispose()


def test_tracker_ten_point_metrics_never_become_rank_evaluations(tmp_path):
    database, importer, workspace = _importer(tmp_path)
    _write(
        workspace / "Application Tracker - Test.md",
        _tracker([("Metrics Co", "Counsel", "**Active — applications open**")]),
    )
    importer.run(workspace)

    with database.session() as session:
        job = session.scalar(select(Job))
        metrics = list(
            session.scalars(select(LegacyMetric).order_by(LegacyMetric.metric_name))
        )
        assert session.scalar(select(func.count()).select_from(Evaluation)) == 0
        assert session.scalar(select(func.count()).select_from(Application)) == 0
        assert score_calibration(session) == []
    assert job.latest_score is None
    assert [
        (metric.metric_name, metric.value, metric.scale_max) for metric in metrics
    ] == [
        ("application_success", 5.0, 10.0),
        ("skills_match", 8.0, 10.0),
    ]
    database.dispose()


@pytest.mark.parametrize(
    ("status", "stage"),
    [
        ("PLANNED — apply in August", ApplicationStage.PLANNED),
        ("✅ SUBMITTED on March 3, 2026", ApplicationStage.APPLIED),
        ("Phone screening scheduled", ApplicationStage.SCREENING),
        ("Interview requested", ApplicationStage.INTERVIEW),
        ("Assessment received", ApplicationStage.ASSESSMENT),
        ("Offer received", ApplicationStage.OFFER),
        ("REJECTED — July 7 (submitted March 3)", ApplicationStage.REJECTED),
        ("WITHDRAWN — application was submitted March 3", ApplicationStage.WITHDRAWN),
        ("ARCHIVED — application sent March 3", ApplicationStage.ARCHIVED),
        ("Applications open — apply now", None),
        ("Application not submitted — draft only", None),
        ("No offer; role remains active", None),
        ("", None),
    ],
)
def test_status_mapping_requires_an_explicit_candidate_application_state(status, stage):
    assert status_to_stage(status) == stage


def test_explicit_tracker_statuses_create_ordered_events_and_advance_on_reimport(
    tmp_path,
):
    database, importer, workspace = _importer(tmp_path)
    tracker = workspace / "Application Tracker - Test.md"
    _write(
        tracker,
        _tracker(
            [
                ("Outcome Co", "Legal Engineer", "✅ **SUBMITTED on March 3, 2026**"),
                ("Open Co", "Policy Fellow", "**Applications open — apply now**"),
            ]
        ),
    )
    importer.run(workspace)

    # The same logical application later has an explicit outcome. Reimport appends
    # an event instead of replacing history or creating another application.
    _write(
        tracker,
        _tracker(
            [
                (
                    "Outcome Co",
                    "Legal Engineer",
                    "❌ **REJECTED — July 7, 2026** (submitted March 3, 2026)",
                ),
                ("Open Co", "Policy Fellow", "**Applications open — apply now**"),
            ]
        ),
    )
    importer.run(workspace)

    with database.session() as session:
        applications = list(session.scalars(select(Application)))
        events = list(
            session.scalars(
                select(StageEvent)
                .where(StageEvent.application_id == applications[0].id)
                .order_by(StageEvent.occurred_at, StageEvent.id)
            )
        )
    assert len(applications) == 1
    assert applications[0].current_stage == ApplicationStage.REJECTED
    assert applications[0].submitted_at.date() == date(2026, 3, 3)
    assert [event.to_stage for event in events] == [
        ApplicationStage.PLANNED,
        ApplicationStage.APPLIED,
        ApplicationStage.REJECTED,
    ]
    assert events[-1].occurred_at.date() == date(2026, 7, 7)
    database.dispose()


def test_pdf_parser_suppresses_only_known_pypdf_warning_level_noise(
    tmp_path, monkeypatch, caplog
):
    class FakeReader:
        def __init__(self, path):
            logger = logging.getLogger("pypdf._reader")
            logger.warning("Ignoring wrong pointing object 1 0 (offset 0)")
            logger.error("A real pypdf error remains visible")
            self.pages = [SimpleNamespace(extract_text=lambda: "PDF text")]

    monkeypatch.setitem(sys.modules, "pypdf", SimpleNamespace(PdfReader=FakeReader))
    caplog.set_level(logging.WARNING)
    pdf = tmp_path / "legacy.pdf"
    pdf.write_bytes(b"fake")

    assert extract_document_text(pdf) == "PDF text"
    assert "Ignoring wrong pointing object" not in caplog.text
    assert "A real pypdf error remains visible" in caplog.text
